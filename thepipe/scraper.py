from typing import Any, Callable, Dict, Iterable, List, Optional, Tuple, Union, cast
import base64
from concurrent.futures import ThreadPoolExecutor, as_completed
from collections import OrderedDict
from io import BytesIO, StringIO
import math
import re
import fnmatch
import os
import tempfile
from urllib.parse import urlparse
import zipfile
from PIL import Image
import requests
import json
import logging

# Module logger (avoid polluting root logger)
logger = logging.getLogger(__name__)
from .drive_utils import extract_drive_id, process_drive_content
from .file_utils import (
    detect_source_type, find_subtitle_files, get_filtered_files, 
    sanitize_filename
)
from .media_utils import (
    MAX_WHISPER_DURATION, VIDEO_PLATFORMS, find_audio_file, 
    clean_subtitles, format_timestamp, get_images_from_markdown,
    get_video_duration_ffmpeg, extract_frame_ffmpeg, extract_audio_segment_ffmpeg,
    # Import the helper functions from media_utils
    process_video, fallback_to_transcription
)
from .web_utils import (
    SCRAPING_PROMPT,
    DRIVE_DOMAINS, GIT_DOMAINS, TWITTER_DOMAINS,
    extract_page_content, matches_domain, normalize_url
)
from .enums import YouTubeEnum
from .core import (
    HOST_IMAGES,
    Chunk,
    make_image_url,
    DEFAULT_AI_MODEL,
    HOST_URL,
)
from .chunker import (
    chunk_by_page,
    chunk_by_document,
    chunk_by_section,
    chunk_semantic,
    chunk_by_keywords,
    chunk_by_length,
    chunk_agentic,
)
import tempfile
import mimetypes
import dotenv
from magika import Magika
import markdownify
from bs4 import BeautifulSoup
import fitz
from openai import OpenAI
from openai.types.chat.chat_completion_message_param import ChatCompletionMessageParam

dotenv.load_dotenv()

FOLDERS_TO_IGNORE = {
    "*node_modules*",
    "*.git*",
    "*venv*",
    "*.vscode*",
    "*pycache*",
    "*.ipynb_checkpoints",
}
FILES_TO_IGNORE = {
    ".gitignore",
    "*.bin",
    # Python compiled files
    "*.pyc",
    "*.pyo",
    "*.pyd",
    # Shared libraries and binaries
    "*.so",
    "*.dll",
    "*.exe",
    # Archives and packages
    "*.tar",
    "*.tar.gz",
    "*.egg-info",
    "package-lock.json",
    "package.json",
    # Lock, log, and metadata files
    "*.lock",
    "*.log",
    "Pipfile.lock",
    "requirements.lock",
    "*.exe",
    "*.dll",
    ".DS_Store",
    "Thumbs.db",
}
GITHUB_TOKEN: Optional[str] = os.getenv("GITHUB_TOKEN", None)
USER_AGENT_STRING: str = os.getenv(
    "USER_AGENT_STRING",
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/58.0.3029.110 Safari/537.3",
)
FILESIZE_LIMIT_MB = int(os.getenv("FILESIZE_LIMIT_MB", 50))

_magika_detector = Magika()

# Global variables for lazy loading
yt_dlp = None

def initialize_video_processing():
    """Initialize video processing libraries."""
    global yt_dlp
    if yt_dlp is None:
        try:
            import yt_dlp
        except ImportError:
            raise ImportError("yt-dlp library not found. Please install it with: pip install yt-dlp")

def is_fifo(path: str) -> bool:
    """Check if a path is a named pipe (FIFO)."""
    import stat
    try:
        return os.path.exists(path) and stat.S_ISFIFO(os.stat(path).st_mode)
    except (OSError, IOError):
        return False


# Module-level constants
COMPILED_BINARY_LABELS = frozenset({
    "executable", "elf_executable", "mach-o", "pe_executable",
    "object", "library", "elf_library", "pe_library", "mach-o_library",
    "java_bytecode", "python_bytecode", "raw_binary", "unknown",
    "archive_executable", "dex", "msi", "nupkg", "deb", "rpm", "apk",
    "font", "firmware", "disk_image", "apple_desktop_services_store",
    "compound_file_binary_format", "lnk", "cat", "mscompress", "cab",
    "pcap", "wasm",
})

# FIFO configuration
FIFO_READ_TIMEOUT = 60  # seconds
FIFO_MAX_SIZE = 100 * 1024 * 1024  # 100MB max


def detect_source_mimetype_from_bytes(data: bytes, filename_hint: Optional[str] = None) -> Optional[str]:
    """Detect MIME type from raw bytes using Magika.
    
    Used for FIFOs and other streams where we can't use identify_path().
    
    Args:
        data: Raw bytes to analyze
        filename_hint: Optional filename for extension-based fallback
    
    Returns:
        MIME type string or None
    """
    try:
        magika_result = _magika_detector.identify_bytes(data)
        
        if magika_result.output.ct_label in COMPILED_BINARY_LABELS:
            return "application/x-compiled-binary"
        
        if magika_result.output.mime_type:
            return magika_result.output.mime_type
        
        # Fallback to extension if we have a filename hint
        if filename_hint:
            guessed_mimetype, _ = mimetypes.guess_type(filename_hint)
            if guessed_mimetype:
                return guessed_mimetype
                
    except Exception as e:
        logger.debug(f"Magika detection failed: {e}")
        if filename_hint:
            guessed_mimetype, _ = mimetypes.guess_type(filename_hint)
            if guessed_mimetype:
                return guessed_mimetype
    
    return None


def _read_fifo_with_timeout(
    filepath: str,
    timeout: int = FIFO_READ_TIMEOUT,
    max_size: int = FIFO_MAX_SIZE,
) -> bytes:
    """Read from a FIFO with timeout and size limit.
    
    Args:
        filepath: Path to the FIFO
        timeout: Maximum seconds to wait for data
        max_size: Maximum bytes to read
        
    Returns:
        Data read from the FIFO
        
    Raises:
        TimeoutError: If no data received within timeout
        ValueError: If data exceeds max_size
        
    Note:
        If timeout occurs, the reader thread may continue running as a daemon.
        This is acceptable since the process will clean up on exit.
    """
    import threading
    
    result = {"data": None, "error": None}
    
    def read_fifo():
        try:
            with open(filepath, 'rb') as f:
                # Read in chunks to enforce size limit
                chunks = []
                total_size = 0
                while True:
                    chunk = f.read(8192)  # 8KB chunks
                    if not chunk:
                        break
                    total_size += len(chunk)
                    if total_size > max_size:
                        result["error"] = ValueError(
                            f"FIFO data exceeds {max_size/(1024*1024):.0f}MB limit"
                        )
                        return
                    chunks.append(chunk)
                result["data"] = b''.join(chunks)
        except Exception as e:
            result["error"] = e
    
    # Run read in thread with timeout
    read_thread = threading.Thread(target=read_fifo, daemon=True)
    read_thread.start()
    read_thread.join(timeout=timeout)
    
    if read_thread.is_alive():
        # Thread is still blocking on open() - FIFO has no writer
        logger.warning(
            f"FIFO read timed out after {timeout}s - reader thread will be orphaned "
            f"(cleanup occurs on process exit)"
        )
        raise TimeoutError(f"FIFO read timed out after {timeout}s - no data received")
    
    if result["error"]:
        raise result["error"]
    
    return result["data"] or b''


def detect_source_mimetype(source: str) -> Optional[str]:
    """Detect the MIME type of a source file, including compiled binaries."""
    if not os.path.isfile(source):
        return None

    _, extension = os.path.splitext(source)
    if extension:
        if extension.lower() == ".ipynb":
            return "application/x-ipynb+json"
        elif extension.lower() in ('.ts', '.tsx'):
            return 'text/'
        elif extension.lower() == '.svg':
            return 'text/'

    try:
        magika_result = _magika_detector.identify_path(source)

        # Use module-level constant (shared with detect_source_mimetype_from_bytes)
        if magika_result.output.ct_label in COMPILED_BINARY_LABELS:
            return "application/x-compiled-binary"
        
        if magika_result.output.mime_type:
            return magika_result.output.mime_type
        
        guessed_mimetype, _ = mimetypes.guess_type(source)
        if guessed_mimetype:
            return guessed_mimetype

    except Exception:
        guessed_mimetype, _ = mimetypes.guess_type(source)
        if guessed_mimetype:
            return guessed_mimetype
        return None

    return None

def scrape_file(
    filepath: str,
    verbose: bool = False,
    chunking_method: Optional[Callable[[List[Chunk]], List[Chunk]]] = None,
    openai_client: Optional[OpenAI] = None,
    model: str = DEFAULT_AI_MODEL,
    text_only: bool = False,
    include_input_images: bool = True,
    include_output_images: bool = True,
    options: Optional[Dict[str, Any]] = None,
) -> List[Chunk]:
    """Scrapes content from various file types, optionally omitting compiled binaries.
    
    Also supports named pipes (FIFOs) as input - will read entire content and detect type.
    """
    if chunking_method is None:
        from .chunker import chunk_by_page
        chunking_method = chunk_by_page

    scraped_chunks = []
    
    # Handle FIFO (named pipe) inputs
    if is_fifo(filepath):
        if verbose:
            print(f"[thepipe] Detected FIFO input: {filepath}")
            print(f"[thepipe] Reading from pipe (timeout: {FIFO_READ_TIMEOUT}s, max: {FIFO_MAX_SIZE/(1024*1024):.0f}MB)...")
        
        try:
            # Read FIFO with timeout and size limit
            fifo_data = _read_fifo_with_timeout(filepath)
            
            if not fifo_data:
                if verbose:
                    print(f"[thepipe] Empty FIFO - no data received")
                return scraped_chunks
            
            if verbose:
                print(f"[thepipe] Read {len(fifo_data)} bytes from FIFO")
            
            # Detect content type from bytes using Magika
            source_mimetype = detect_source_mimetype_from_bytes(fifo_data)
            
            if source_mimetype is None:
                if verbose:
                    print(f"[thepipe] Could not detect content type from FIFO data")
                return scraped_chunks
            
            if verbose:
                print(f"[thepipe] Detected FIFO content type: {source_mimetype}")
            
            # Write to temp file so existing scrapers can process it
            import tempfile
            suffix = mimetypes.guess_extension(source_mimetype) or ''
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                tmp.write(fifo_data)
                temp_filepath = tmp.name
            
            try:
                # Recursively call scrape_file with the temp file
                scraped_chunks = scrape_file(
                    filepath=temp_filepath,
                    verbose=verbose,
                    chunking_method=chunking_method,
                    openai_client=openai_client,
                    model=model,
                    text_only=text_only,
                    include_input_images=include_input_images,
                    include_output_images=include_output_images,
                    options=options,
                )
                # Update path references to point to original FIFO
                for chunk in scraped_chunks:
                    if chunk.path == temp_filepath:
                        chunk.path = filepath
            finally:
                # Cleanup temp file
                os.unlink(temp_filepath)
            
            return scraped_chunks
            
        except Exception as e:
            if verbose:
                print(f"[thepipe] Error reading from FIFO: {e}")
            return scraped_chunks
    
    source_mimetype = detect_source_mimetype(filepath)

    # Determine if compiled binaries should be read
    # Defaults to False, can be overridden by options["read_executable"]
    read_executable = options.get("read_executable", False) if options else False

    if source_mimetype is None:
        if verbose:
            print(f"[thepipe] Unsupported or unreadable source: {filepath}")
        return scraped_chunks
        
    # Skip compiled binaries unless read_executable is explicitly True
    if source_mimetype == "application/x-compiled-binary" and not read_executable:
        if verbose:
            print(f"[thepipe] Skipping detected compiled binary file: {filepath}")
        return scraped_chunks

    if verbose:
        print(f"[thepipe] Scraping {source_mimetype}: {filepath}...")

    if source_mimetype == "application/pdf":
        scraped_chunks = scrape_pdf(
            file_path=filepath, verbose=verbose, model=model, options=options,
            openai_client=openai_client, include_input_images=include_input_images,
            include_output_images=include_output_images
        )
    elif source_mimetype == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        scraped_chunks = scrape_docx(
            file_path=filepath, verbose=verbose, include_output_images=include_output_images,
            options=options
        )
    elif source_mimetype == "application/vnd.openxmlformats-officedocument.presentationml.presentation":
        scraped_chunks = scrape_pptx(
            file_path=filepath, verbose=verbose, include_output_images=include_output_images,
            options=options
        )
    elif source_mimetype.startswith("image/"):
        scraped_chunks = scrape_image(file_path=filepath, options=options)
    elif source_mimetype.startswith("application/vnd.ms-excel") or \
         source_mimetype == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet":
        scraped_chunks = scrape_spreadsheet(file_path=filepath, source_type=source_mimetype, options=options)
    elif source_mimetype == "application/x-ipynb+json":
        scraped_chunks = scrape_ipynb(
            file_path=filepath, verbose=verbose, include_output_images=include_output_images
        )
    elif source_mimetype == "application/zip" or source_mimetype == "application/x-zip-compressed":
        scraped_chunks = scrape_zip(
            file_path=filepath, verbose=verbose, text_only=text_only, options=options,
            openai_client=openai_client, include_input_images=include_input_images,
            include_output_images=include_output_images,
        )
    elif source_mimetype.startswith("video/"):
        scraped_chunks = scrape_video(
            file_path=filepath, verbose=verbose, include_output_images=include_output_images,
            text_only=text_only, options=options
        )
    elif source_mimetype.startswith("audio/"):
        scraped_chunks = scrape_audio(file_path=filepath, verbose=verbose, options=options)
    elif source_mimetype.startswith("text/html"):
        scraped_chunks = scrape_html(
            file_path=filepath, verbose=verbose, include_output_images=include_output_images,
            options=options
        )
    elif source_mimetype.startswith("text/"):
        scraped_chunks = scrape_plaintext(file_path=filepath, options=options)
    else:
        try:
            scraped_chunks = scrape_plaintext(file_path=filepath, options=options)
        except Exception as e:
            if verbose:
                print(f"[thepipe] Error extracting from {filepath} (fallback to plaintext): {e}")
                
    if verbose:
        if scraped_chunks:
            print(f"[thepipe] Extracted from {filepath}")
        else:
            print(f"[thepipe] No content extracted from {filepath}")
            
    if chunking_method:
        scraped_chunks = chunking_method(scraped_chunks)
        
    return scraped_chunks

def scrape_html(
    file_path: str,
    verbose: bool = False,
    include_output_images: bool = True,
    options: Optional[Dict[str, Any]] = None
) -> List[Chunk]:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
        html_content = file.read()

    # Parse HTML and remove unwanted elements for cleaner content
    soup = BeautifulSoup(html_content, "html.parser")
    for script in soup(["script", "style", "nav", "footer", "header"]):
        script.decompose()

    markdown_content = markdownify.markdownify(str(soup), heading_style="ATX")
    images = get_images_from_markdown(html_content) if include_output_images else []
    return [Chunk(path=file_path, text=markdown_content, images=images)]

def scrape_plaintext(file_path: str, options: Optional[Dict[str, Any]] = None ) -> List[Chunk]:
    with open(file_path, "r", encoding="utf-8", errors="ignore") as file:
        text = file.read()
    return [Chunk(path=file_path, text=text)]

def scrape_directory(
    dir_path: str,
    include_regex: Optional[str] = None,
    include_patterns: Optional[List[str]] = None,
    verbose: bool = False,
    text_only: bool = False,
    options: Optional[Dict[str, Any]] = None,
    openai_client: Optional[OpenAI] = None,
    model: str = DEFAULT_AI_MODEL,
    include_input_images: bool = True,
    include_output_images: bool = True,
) -> List[Chunk]:
    extraction = []
    
    # Process options
    options = options or {}
    
    # Check for code_relations mode
    # Options: "limited", "map", "mapnn", "mapall"
    # Additional: code_n1 (default 3), code_n2 (default 5)
    code_relations = options.get('code_relations')
    if code_relations:
        from .analyzer.integration import process_code_relations
        return process_code_relations(
            dir_path=dir_path,
            include_patterns=include_patterns,
            mode=code_relations,
            code_n1=options.get('code_n1', 3),
            code_n2=options.get('code_n2', 5),
            code_nf=options.get('code_nf', 100),
            code_nt=options.get('code_nt', 150000),
            verbose=verbose,
            options=options,
        )
    
    # Get blacklist files list (if any)
    blacklist_files = options.get('blacklist_files', [])
    
    # Get filtered file list
    all_files = get_filtered_files(
        dir_path=dir_path,
        include_regex=include_regex,
        include_patterns=include_patterns,
        blacklist_files=blacklist_files,
        verbose=verbose
    )
    
    if verbose:
        print(f"[thepipe] Processing {len(all_files)} files in {dir_path}")
    
    with ThreadPoolExecutor() as executor:
        results = executor.map(
            lambda file_path: scrape_file(
                filepath=file_path,
                model=model,
                verbose=verbose,
                options=options,
                openai_client=openai_client,
                include_input_images=include_input_images,
                include_output_images=include_output_images,
            ),
            all_files,
        )
        for result in results:
            extraction.extend(result)
    
    return extraction

def scrape_zip(
    file_path: str,
    include_regex: Optional[str] = None,
    include_patterns: Optional[List[str]] = None,
    verbose: bool = False,
    text_only: bool = False,
    options: Optional[Dict[str, Any]] = None,
    openai_client: Optional[OpenAI] = None,
    include_input_images: bool = True,
    include_output_images: bool = True,
) -> List[Chunk]:
    chunks = []
    with tempfile.TemporaryDirectory() as temp_dir:
        with zipfile.ZipFile(file_path, "r") as zip_ref:
            zip_ref.extractall(temp_dir)
        chunks = scrape_directory(
            dir_path=temp_dir,
            include_regex=include_regex,
            include_patterns=include_patterns,
            verbose=verbose,
            text_only=text_only,
            options=options,
            openai_client=openai_client,
            include_input_images=include_input_images,
            include_output_images=include_output_images,
        )
    return chunks

def scrape_pdf(
    file_path: str,
    model: str = DEFAULT_AI_MODEL,
    verbose: Optional[bool] = False,
    options: Optional[Dict[str, Any]] = None,
    openai_client: Optional[OpenAI] = None,
    include_input_images: bool = True,
    include_output_images: bool = True,
    image_scale: float = 1.0,
) -> List[Chunk]:
    chunks: List[Chunk] = []
    MAX_PAGES = 128

    # Branch 1 – VLM path (AI extraction or OpenAI client supplied)
    if openai_client is not None:
        from collections import OrderedDict
        import concurrent.futures

        if openai_client is None:
            # Create client from options or environment
            llm_config = options.get("llm_extractor", {}) if options else {}
            api_key = llm_config.get("api_key", os.environ.get("LLM_SERVER_API_KEY", os.environ.get("OPENAI_API_KEY")))
            api_base = llm_config.get("api_base", os.environ.get("LLM_SERVER_BASE_URL"))
            model = llm_config.get("model", model or DEFAULT_AI_MODEL)
            
            # Configure OpenAI client
            client_args = {"api_key": api_key}
            if api_base:
                client_args["base_url"] = api_base
                
            openai_client = OpenAI(**client_args)

        with open(file_path, "rb") as fp:
            pdf_bytes = fp.read()
        doc = fitz.open(stream=pdf_bytes, filetype="pdf")
        num_pages = len(doc)

        if num_pages > MAX_PAGES:
            err = f"Error: PDF has {num_pages} pages (max is {MAX_PAGES} for AI extraction)."
            raise Exception(err)

        if verbose:
            print(
                f"[thepipe] Scraping PDF: {file_path} "
                f"({num_pages} pages) with model {model}"
            )

        # Inner worker – processes one page
        def _process_page(page_num: int) -> Tuple[int, str, Optional[Image.Image]]:
            page = doc[page_num]
            text = page.get_text()  # type: ignore[attr-defined]

            # Build message for the LLM
            msg_content: List[Dict[str, Union[Dict[str, str], str]]] = [
                {
                    "type": "text",
                    "text": f"```\n{text}\n```\n{SCRAPING_PROMPT}",
                }
            ]

            image: Optional[Image.Image] = None
            if include_input_images or include_output_images:
                mat = fitz.Matrix(image_scale, image_scale)
                pix = page.get_pixmap(matrix=mat, alpha=False)  # type: ignore[attr-defined]  # noqa: E501
                image = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)

                if include_input_images:
                    encoded = make_image_url(image, host_images=HOST_IMAGES)
                    msg_content.append(
                        {
                            "type": "image_url",
                            "image_url": {"url": encoded, "detail": "high"},
                        }
                    )

            messages = [{"role": "user", "content": msg_content}]

            # Use unified LLMClient (supports both OpenAI and agent mode)
            from .llm import LLMClient
            llm_client = LLMClient.from_options(
                options=options,
                model=model,
            )
            response = llm_client.query(messages=messages)

            llm_response = response.content
            if not llm_response:
                raise RuntimeError("Empty LLM response.")

            llm_response = llm_response.strip()
            if llm_response.startswith("```markdown"):
                llm_response = llm_response[len("```markdown") :]
            elif llm_response.startswith("```"):
                llm_response = llm_response[len("```") :]
            if llm_response.endswith("```"):
                llm_response = llm_response[: -len("```")]

            return (
                page_num,
                llm_response,
                image if include_output_images else None,
            )

        # Parallel extraction
        max_workers = (os.cpu_count() or 1) * 2
        if verbose:
            print(f"[thepipe] Using {max_workers} threads for PDF extraction")

        page_results: OrderedDict[int, Tuple[str, Optional[Image.Image]]] = (
            OrderedDict()
        )
        with ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = [executor.submit(_process_page, p) for p in range(num_pages)]
            for fut in as_completed(futures):
                pg, txt, img = fut.result()
                page_results[pg] = (txt, img)

        for pg in sorted(page_results):
            txt, img = page_results[pg]
            chunks.append(Chunk(path=file_path, text=txt, images=[img] if img else []))

        return chunks

    # Branch 2 – no OpenAI client – text-only offline mode
    from pymupdf4llm.helpers.pymupdf_rag import to_markdown  # local import

    doc = fitz.open(file_path)
    md_pages = cast(List[Dict[str, Any]], to_markdown(file_path, page_chunks=True))

    for i in range(doc.page_count):
        text = re.sub(r"\n{3,}", "\n\n", md_pages[i]["text"]).strip()

        images: List[Image.Image] = []
        if include_output_images:
            mat = fitz.Matrix(image_scale, image_scale)
            pix = doc[i].get_pixmap(matrix=mat, alpha=False)  # type: ignore[attr-defined]  # noqa: E501
            images.append(Image.frombytes("RGB", [pix.width, pix.height], pix.samples))

        chunks.append(Chunk(path=file_path, text=text, images=images))

    doc.close()
    return chunks

def scrape_image(file_path: str, options: Optional[Dict[str, Any]] = None ) -> List[Chunk]:
    img = Image.open(file_path)
    img.load()  # needed to close the file
    chunk = Chunk(path=file_path, images=[img])
    return [chunk]

def scrape_spreadsheet(file_path: str, source_type: str, options: Optional[Dict[str, Any]] = None) -> List[Chunk]:
    import pandas as pd

    if source_type == "application/vnd.ms-excel":
        df = pd.read_csv(file_path)
    elif (
        source_type
        == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ):
        df = pd.read_excel(file_path)
    else:
        raise ValueError("Unsupported file format")
    dicts = df.to_dict(orient="records")
    chunks = []
    for i, item in enumerate(dicts):
        # format each row as json along with the row index
        item["row index"] = i
        item_json = json.dumps(item, indent=4)
        chunks.append(Chunk(path=file_path, text=item_json))
    return chunks

def scrape_url(
    url: str, 
    include_regex: Optional[str] = None,
    include_patterns: Optional[List[str]] = None, 
    text_only: bool = False, 
    verbose: bool = False, 
    chunking_method: Optional[Callable] = chunk_by_page, 
    options: Optional[Dict[str, Any]] = None,
    openai_client: Optional[OpenAI] = None,
    model: str = DEFAULT_AI_MODEL,
    include_input_images: bool = True,
    include_output_images: bool = True,
) -> List[Chunk]:
    """Scrape content from a URL."""
    cookie_options = options.get('cookies', {}) if options else {}
    
    # Handle cookie test mode early
    if cookie_options.get('show') == "test" and cookie_options.get('to_terminal', True):
        from .cookie_utils import process_cookie_options
        cookie_info = process_cookie_options(url, [], cookie_options)
        if isinstance(cookie_info, str):
            print(cookie_info)
            return []
        return cookie_info

    # Process URL locally
    chunks = []
    try:
        if matches_domain(url, DRIVE_DOMAINS):
            if verbose:
                print("[thepipe] Detected Google Drive/Docs URL, using drive scraper")
            chunks = scrape_drive(url, text_only=text_only,
                                verbose=verbose, options=options)
        elif matches_domain(url, VIDEO_PLATFORMS):
            chunks = scrape_youtube(url, text_only=text_only,
                                  verbose=verbose, options=options)
        elif matches_domain(url, TWITTER_DOMAINS):
            chunks = scrape_tweet(url=url, text_only=text_only,
                                verbose=verbose, options=options)
        elif matches_domain(url, GIT_DOMAINS):
            chunks = scrape_github(github_url=url, include_regex=include_regex,
                                 include_patterns=include_patterns,
                                 text_only=text_only, verbose=verbose, options=options)
        else:
            # Handle other content types
            parsed_url = urlparse(normalize_url(url))
            file_extension = os.path.splitext(parsed_url.path)[1].lower()
            if file_extension in ['.pdf', '.docx', '.txt', '.csv', '.xlsx']:
                with tempfile.TemporaryDirectory() as temp_dir:
                    file_path = os.path.join(temp_dir, os.path.basename(parsed_url.path))
                    response = requests.get(normalize_url(url))
                    if (FILESIZE_LIMIT_MB and 
                        int(response.headers.get('Content-Length', 0)) > FILESIZE_LIMIT_MB * 1024 * 1024):
                        raise ValueError(f"File size exceeds {FILESIZE_LIMIT_MB} MB limit.")
                    with open(file_path, 'wb') as file:
                        file.write(response.content)
                    chunks = scrape_file(
                        filepath=file_path, 
                        text_only=text_only, 
                        verbose=verbose,
                        chunking_method=chunking_method,
                        options=options,
                        openai_client=openai_client,
                        include_input_images=include_input_images,
                        include_output_images=include_output_images,
                    )
            else:
                if openai_client and include_input_images:
                    chunk = parse_webpage_with_vlm(
                        url=url,
                        verbose=verbose,
                        model=model,
                        openai_client=openai_client,
                        include_output_images=include_output_images,
                    )
                else:
                    chunk = extract_page_content(
                        url=url, 
                        verbose=verbose, 
                        include_output_images=include_output_images,
                    )
                chunks = chunking_method([chunk])
                if not any(chunk.text for chunk in chunks) and not any(chunk.images for chunk in chunks):
                    raise ValueError("No content extracted from URL.")
                    
    except ImportError as e:
        raise ImportError(f"Required dependencies not found: {str(e)}")
    except Exception as e:
        if verbose:
            print(f"[thepipe] Error processing URL: {str(e)}")
        raise

    # Process any cookie options if present
    if cookie_options:
        from .cookie_utils import process_cookie_options
        return process_cookie_options(url, chunks, cookie_options)
    return chunks

def scrape_drive(drive_url: str, text_only: bool = False, 
                verbose: bool = False, 
                options: Optional[Dict[str, Any]] = None) -> List[Chunk]:
    """Process Google Drive URLs (both files and folders)."""
    if verbose:
        print(f"[thepipe] Processing Drive URL: {drive_url}")

    drive_id = extract_drive_id(drive_url)
    if not drive_id:
        raise ValueError(f"Could not extract Drive ID from URL: {drive_url}")
        
    return process_drive_content(
        drive_url=drive_url,
        drive_id=drive_id,
        text_only=text_only,
        verbose=verbose,
        options=options
    )

def parse_webpage_with_vlm(
    url: str,
    model: str = DEFAULT_AI_MODEL,
    verbose: Optional[bool] = False,
    openai_client: Optional[OpenAI] = None,  # Legacy, kept for compatibility
    include_output_images: bool = True,
    options: Optional[Dict[str, Any]] = None,  # New: options with llm_provider
) -> Chunk:
    # LLMClient handles agent mode automatically via options
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = p.chromium.launch()
        context = browser.new_context(user_agent=USER_AGENT_STRING)
        page = context.new_page()
        page.goto(url, wait_until="domcontentloaded")
        if not page.viewport_size:
            page.set_viewport_size({"width": 800, "height": 600})
        if not page.viewport_size:
            raise ValueError(
                "Failed to set viewport size after finding no viewport size"
            )
        viewport_height = page.viewport_size.get("height", 800)
        total_height = page.evaluate("document.body.scrollHeight")
        current_scroll_position = 0
        scrolldowns, max_scrolldowns = 0, 3
        images: List[Image.Image] = []

        while current_scroll_position < total_height and scrolldowns < max_scrolldowns:
            page.wait_for_timeout(200)  # wait for content to load
            screenshot = page.screenshot(full_page=False)
            img = Image.open(BytesIO(screenshot))
            images.append(img)

            current_scroll_position += viewport_height
            page.evaluate(f"window.scrollTo(0, {current_scroll_position})")
            scrolldowns += 1
            total_height = page.evaluate("document.body.scrollHeight")
            if verbose:
                print(
                    f"[thepipe] Scrolled to {current_scroll_position} of {total_height}. Waiting for content to load..."
                )

        browser.close()

    if images:
        # Vertically stack the images
        total_height = sum(img.height for img in images)
        max_width = max(img.width for img in images)
        stacked_image = Image.new("RGB", (max_width, total_height))
        y_offset = 0
        for img in images:
            stacked_image.paste(img, (0, y_offset))
            y_offset += img.height

        # Process the stacked image with VLM
        messages = [
            {
                "role": "user",
                "content": [
                    {
                        "type": "image_url",
                        "image_url": {
                            "url": make_image_url(
                                stacked_image, host_images=HOST_IMAGES
                            ),
                            "detail": "high",
                        },
                    },
                    {"type": "text", "text": SCRAPING_PROMPT},
                ],
            },
        ]
        # Use unified LLMClient (supports both OpenAI and agent mode)
        from .llm import LLMClient
        llm_client = LLMClient.from_options(
            options=options,
            model=model,
        )
        response = llm_client.query(messages=messages)
        llm_response = response.content
        if not llm_response:
            raise Exception(
                f"Failed to receive a message content from LLM Response"
            )
        if verbose:
            print(f"[thepipe] LLM response: {llm_response}")
        chunk = Chunk(
            path=url,
            text=llm_response,
            images=[stacked_image] if include_output_images else [],
        )
    else:
        raise ValueError("Model received 0 images from webpage")

    return chunk

def extract_page_content(
    url: str, verbose: bool = False, include_output_images: bool = True
) -> Chunk:
    from bs4 import BeautifulSoup
    from playwright.sync_api import sync_playwright
    import base64
    import requests

    texts: List[str] = []
    images: List[Image.Image] = []

    with sync_playwright() as p:
        browser = p.chromium.launch()
        try:
            context = browser.new_context(user_agent="USER_AGENT_STRING")
            page = context.new_page()
            page.goto(url, wait_until="domcontentloaded")

            # Scroll to the bottom of the page to load dynamic content
            if not page.viewport_size:
                page.set_viewport_size({"width": 800, "height": 600})
            if not page.viewport_size:
                raise ValueError(
                    "Failed to set viewport size after finding no viewport size"
                )
            viewport_height = page.viewport_size["height"]
            total_height = page.evaluate("document.body.scrollHeight")
            current_scroll_position = 0
            scrolldowns, max_scrolldowns = 0, 20  # Finite to prevent infinite scroll

            while current_scroll_position < total_height and scrolldowns < max_scrolldowns:
                page.wait_for_timeout(200)  # Wait for dynamic content to load
                current_scroll_position += viewport_height
                page.evaluate(f"window.scrollTo(0, {current_scroll_position})")
                scrolldowns += 1
                total_height = page.evaluate("document.body.scrollHeight")

            # Extract HTML content
            html_content = page.content()

            # Convert HTML to Markdown
            soup = BeautifulSoup(html_content, "html.parser")

            # Remove script, style, and navigation elements for cleaner content
            for script in soup(["script", "style", "nav", "footer", "header"]):
                script.decompose()

            markdown_content = markdownify.markdownify(str(soup), heading_style="ATX")

            # Remove excessive newlines in the markdown
            markdown_content = re.sub(r"\n{3,}", "\n\n", markdown_content)
            markdown_content = markdown_content.strip()

            texts.append(markdown_content)

            # Extract images from the page using heuristics
            if include_output_images:
                for img in page.query_selector_all("img"):
                    img_path = img.get_attribute("src")
                    if not img_path:
                        continue
                    if img_path.startswith("data:image"):
                        # Save base64 image to PIL Image
                        decoded_data = base64.b64decode(img_path.split(",")[1])
                        try:
                            image = Image.open(BytesIO(decoded_data))
                            images.append(image)
                        except Exception as e:
                            if verbose:
                                print(
                                    f"[thepipe] Ignoring error loading image {img_path}: {e}"
                                )
                            continue  # Ignore incompatible image extractions
                    else:
                        try:
                            response = requests.get(
                                img_path,
                                timeout=10,
                                headers={"User-Agent": USER_AGENT_STRING},
                            )
                            response.raise_for_status()
                            image = Image.open(BytesIO(response.content))
                            images.append(image)
                        except Exception as e:
                            if verbose:
                                print(
                                    f"[thepipe] Ignoring error loading image {img_path}: {e}"
                                    "Attempting to load path with schema."
                                )

                            if "https://" not in img_path and "http://" not in img_path:
                                try:
                                    while img_path.startswith("/"):
                                        img_path = img_path[1:]
                                    path_with_schema = (
                                        urlparse(url).scheme + "://" + img_path
                                    )
                                    response = requests.get(
                                        path_with_schema,
                                        timeout=10,
                                        headers={"User-Agent": USER_AGENT_STRING},
                                    )
                                    response.raise_for_status()
                                    image = Image.open(BytesIO(response.content))
                                    images.append(image)
                                except Exception as e:
                                    if verbose:
                                        print(
                                            f"[thepipe] Ignoring error loading image {img_path} with schema: {e}"
                                            "Attempting to load with schema and netloc."
                                        )
                                    try:
                                        path_with_schema_and_netloc = (
                                            urlparse(url).scheme
                                            + "://"
                                            + urlparse(url).netloc
                                            + "/"
                                            + img_path
                                        )
                                        response = requests.get(
                                            path_with_schema_and_netloc,
                                            timeout=10,
                                            headers={"User-Agent": USER_AGENT_STRING},
                                        )
                                        response.raise_for_status()
                                        image = Image.open(BytesIO(response.content))
                                        images.append(image)
                                    except Exception as e:
                                        if verbose:
                                            print(
                                                f"[thepipe] Ignoring error loading image {img_path} with schema and netloc: {e}"
                                                "Continuing."
                                            )
                                        continue  # Ignore incompatible image extractions
                            else:
                                if verbose:
                                    print(
                                        f"[thepipe] Ignoring error loading image {img_path}"
                                    )
                                continue  # Ignore incompatible image extractions
        except Exception as e:
            if verbose:
                print(f"[thepipe] Error scraping {url}: {e}")
        finally:
            browser.close()

    text = "\n".join(texts).strip()
    return Chunk(path=url, text=text, images=images)

def scrape_video(file_path: str, verbose: bool = False, include_output_images: bool = True, 
                text_only: Optional[Union[bool, str]] = None, options: Optional[Dict[str, Any]] = None) -> List[Chunk]:
    """Scrape video content using FFmpeg for better performance and compatibility."""
    try:
        import whisper
    except ImportError:
        raise ImportError("whisper library not found. Please install it with: pip install openai-whisper")
    
    model = whisper.load_model("base")
    chunks = []
    
    try:
        # Get video duration using FFmpeg
        duration = get_video_duration_ffmpeg(file_path)
        
        if verbose:
            print(f"[thepipe] Video duration: {duration:.1f} seconds")
        
        # Calculate number of chunks
        num_chunks = math.ceil(duration / MAX_WHISPER_DURATION)
        
        for i in range(num_chunks):
            start_time = i * MAX_WHISPER_DURATION
            end_time = min(start_time + MAX_WHISPER_DURATION, duration)
            chunk_duration = end_time - start_time
            
            if verbose:
                print(f"[thepipe] Processing chunk {i+1}/{num_chunks} ({start_time:.1f}s - {end_time:.1f}s)")
            
            # Extract frame from middle of chunk (if not text_only)
            image = None
            if not text_only:
                frame_time = (start_time + end_time) / 2
                image = extract_frame_ffmpeg(file_path, frame_time, verbose)
            
            # Extract and transcribe audio segment
            transcription = None
            audio_path = None
            try:
                # Create temporary audio file
                with tempfile.NamedTemporaryFile(suffix=".wav", delete=False) as temp_audio:
                    audio_path = temp_audio.name
                
                # Extract audio segment using FFmpeg
                if extract_audio_segment_ffmpeg(file_path, start_time, chunk_duration, audio_path, verbose):
                    # Transcribe with Whisper
                    result = model.transcribe(audio=audio_path, verbose=verbose)
                    
                    # Format transcription with global timestamps
                    formatted_transcription = []
                    for segment in result["segments"]:
                        # Adjust timestamps to global video time
                        global_start = start_time + segment["start"]
                        global_end = start_time + segment["end"]
                        start_ts = format_timestamp(global_start, 0, 0)
                        end_ts = format_timestamp(global_end, 0, 0)
                        formatted_transcription.append(
                            f"[{start_ts} --> {end_ts}]  {segment['text']}"
                        )
                    transcription = "\n".join(formatted_transcription)
                else:
                    if verbose:
                        print(f"[thepipe] Failed to extract audio for chunk {i+1}")
                        
            except Exception as e:
                if verbose:
                    print(f"[thepipe] Error processing audio for chunk {i+1}: {str(e)}")
            finally:
                # Clean up temporary audio file
                if audio_path and os.path.exists(audio_path):
                    try:
                        os.remove(audio_path)
                    except Exception as e:
                        if verbose:
                            print(f"[thepipe] Warning: Could not remove temp file {audio_path}: {e}")
            
            # Create chunk with available content
            texts = [transcription] if transcription else []
            images = [image] if image and not text_only else []
            
            if texts or images:
                chunks.append(Chunk(path=file_path, text=texts[0] if texts else "", images=images))
            elif verbose:
                print(f"[thepipe] No content extracted for chunk {i+1}")
    
    except Exception as e:
        if verbose:
            print(f"[thepipe] Error processing video {file_path}: {str(e)}")
        # Return error chunk
        chunks.append(Chunk(path=file_path, text=f"Error processing video: {str(e)}"))
    
    return chunks

def scrape_youtube(url: str, text_only: Optional[Union[bool, str]] = None, verbose: bool = False, 
                   metadata_fields: Optional[List[YouTubeEnum]] = None, 
                   options: Optional[Dict[str, Any]] = None) -> List[Chunk]:
    """Scrape content from a YouTube URL with robust error handling and fallback to transcription."""
    initialize_video_processing()
    if verbose:
        print("[thepipe] Initializing YouTube content extraction...")
    
    ydl_opts = {
        'quiet': not verbose,
        'ignoreerrors': True,
        'extract_flat': 'in_playlist',
        'outtmpl': '%(title)s.%(ext)s',
        'subtitlesformat': 'vtt',
        'skip_download': True,  # Always skip video download initially
    }

    # Configure subtitle extraction
    if text_only != 'transcribe':
        ydl_opts.update({
            'writesubtitles': True,
            'writeautomaticsub': True,
            'subtitleslangs': ['en', 'en-orig'],
        })

    if options and 'youtube' in options:
        ydl_opts.update(YouTubeEnum.process_options(options['youtube'], bool(text_only), verbose))

    chunks = []
    with tempfile.TemporaryDirectory() as temp_dir:
        ydl_opts['outtmpl'] = os.path.join(temp_dir, '%(title)s.%(ext)s')

        try:
            with yt_dlp.YoutubeDL(ydl_opts) as ydl:
                # CRITICAL FIX: Validate extract_info response before processing
                info = ydl.extract_info(url, download=False)
                
                if not isinstance(info, dict):
                    if verbose:
                        print(f"[thepipe] Invalid info type from extract_info: {type(info)}")
                    # Try to get basic video info and fall back to transcription
                    return fallback_to_transcription(url, temp_dir, verbose)
                
                if 'entries' in info:  # It's a playlist
                    videos = info['entries']
                else:  # Single video
                    videos = [info]

                for video in videos:
                    if video is None:
                        if verbose:
                            print("[thepipe] Skipping None video entry")
                        continue
                        
                    # ADDITIONAL FIX: Validate each video entry
                    if not isinstance(video, dict):
                        if verbose:
                            print(f"[thepipe] Invalid video entry type: {type(video)}")
                        continue
                        
                    video_chunks = process_video(ydl, video, temp_dir, text_only, verbose, metadata_fields)
                    chunks.extend(video_chunks)

        except Exception as e:
            if verbose:
                print(f"[thepipe] Error processing content: {str(e)}")
            # Try fallback to transcription
            fallback_chunks = fallback_to_transcription(url, temp_dir, verbose)
            if fallback_chunks:
                chunks.extend(fallback_chunks)
            else:
                chunks.append(Chunk(path=url, text=f"Error: Unable to process content. {str(e)}"))

    return chunks

def scrape_audio(file_path: str, verbose: bool = False, options: Optional[Dict[str, Any]] = None) -> List[Chunk]:
    """Transcribe audio file using Whisper with robust error handling."""
    try:
        import whisper
    except ImportError:
        raise ImportError("whisper library not found. Please install it with: pip install openai-whisper")

    try:
        model = whisper.load_model("base")
        if verbose:
            print(f"[thepipe] Transcribing audio file: {file_path}")
        
        result = model.transcribe(audio=file_path, verbose=verbose)
        segments = result.get("segments", [])

        transcript: List[str] = []
        for segment in segments:
            start = format_timestamp(segment["start"], 0, 0)
            end = format_timestamp(segment["end"], 0, 0)
            if segment["text"].strip():
                transcript.append(f"[{start} --> {end}]  {segment['text']}")
        
        # Join the formatted transcription into a single string
        transcription_text = '\n'.join(transcript)
        if verbose:
            print(f"[thepipe] Transcription completed for {file_path}")
        return [Chunk(path=file_path, text=transcription_text)]
        
    except Exception as e:
        if verbose:
            print(f"[thepipe] Error transcribing audio: {str(e)}")
        return [Chunk(path=file_path, text=f"Error transcribing audio: {str(e)}")]

def scrape_github(
    github_url: str, include_regex: Optional[str] = None, include_patterns: Optional[List[str]] = None,
    text_only: bool = False, branch: str = "main", verbose: bool = False,
    options: Optional[Dict[str, Any]] = None) -> List[Chunk]:
    """Scrape content from a GitHub repository with optional authentication."""
    with tempfile.TemporaryDirectory() as temp_dir:
        # Try unauthenticated clone first
        clone_result = os.system(f"git clone {github_url} {temp_dir} --quiet")
        
        # If clone fails and we have token options/env, try authenticated clone
        if clone_result != 0:
            # Check options first, then environment variable
            token = None
            if options:
                token = options.get('github_token') or options.get('github', {}).get('token')
            if not token:
                token = os.getenv('GITHUB_TOKEN')
                
            if token:
                if verbose:
                    print(f"[thepipe] Attempting authenticated clone...")
                auth_url = github_url.replace("https://", f"https://{token}@")
                clone_result = os.system(f"git clone {auth_url} {temp_dir} --quiet")
                if clone_result != 0:
                    return [Chunk(path=github_url, text=f"Failed to clone repository even with authentication: {github_url}")]
            else:
                return [Chunk(path=github_url, text=f"Repository requires authentication. Set GITHUB_TOKEN environment variable or provide token in options")]

        try:
            blacklist_files = options.get('blacklist_files', []) if options else []
            if options and options.get('gitignore', False):
                git_ignore_files = ['.gitignore', '.git/info/exclude']
                for git_file in git_ignore_files:
                    if os.path.exists(os.path.join(temp_dir, git_file)) and git_file not in blacklist_files:
                        blacklist_files.append(git_file)
                        if verbose:
                            print(f"[thepipe] Adding {git_file} to blacklist files")
            
            modified_options = (options or {}).copy()
            modified_options['blacklist_files'] = blacklist_files
            
            if verbose and blacklist_files:
                print(f"[thepipe] Using blacklist files: {blacklist_files}")
                
            return scrape_directory(
                dir_path=temp_dir,
                include_regex=include_regex,
                include_patterns=include_patterns,
                verbose=verbose,
                text_only=text_only,
                options=modified_options
            )
        except Exception as e:
            if verbose:
                print(f"[thepipe] Error processing repository contents: {str(e)}")
            return [Chunk(path=github_url, text=f"Error processing repository contents: {str(e)}")]

def scrape_docx(
    file_path: str,
    verbose: bool = False,
    include_output_images: bool = True,
    options: Optional[Dict[str, Any]] = None
) -> List[Chunk]:
    from docx import Document
    from docx.oxml.table import CT_Tbl
    from docx.oxml.text.paragraph import CT_P
    from docx.table import Table, _Cell
    from docx.text.paragraph import Paragraph
    import csv
    import io

    # helper function to iterate through blocks in the document
    def iter_block_items(parent):
        if parent.__class__.__name__ == "Document":
            parent_elm = parent.element.body
        elif parent.__class__.__name__ == "_Cell":
            parent_elm = parent._tc
        else:
            raise ValueError("Unsupported parent type")
        # iterate through each child element in the parent element
        for child in parent_elm.iterchildren():
            child_elem_class_name = child.__class__.__name__
            if verbose:
                print(f"[thepipe] Found element in docx: {child_elem_class_name}")
            if child_elem_class_name == "CT_P":
                yield Paragraph(child, parent)
            elif child_elem_class_name == "CT_Tbl":
                yield Table(child, parent)

    # helper function to read tables in the document
    def read_docx_tables(tab):
        vf = StringIO()
        writer = csv.writer(vf)
        for row in tab.rows:
            writer.writerow(cell.text for cell in row.cells)
        vf.seek(0)
        return vf.getvalue()

    # read the document
    document = Document(file_path)
    chunks = []
    image_counter = 0

    # Define namespaces
    nsmap = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
        "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    }

    try:
        # scrape each block in the document to create chunks
        for block in iter_block_items(document):
            block_texts = []
            block_images = []
            if isinstance(block, Paragraph):
                block_texts.append(block.text)
                # "runs" are the smallest units in a paragraph
                for run in block.runs:
                    if "pic:pic" in run.element.xml and include_output_images:
                        # extract images from the paragraph
                        for pic in run.element.findall(".//pic:pic", nsmap):
                            cNvPr = pic.find(".//pic:cNvPr", nsmap)
                            name_attr = (
                                cNvPr.get("name")
                                if cNvPr is not None
                                else f"image_{image_counter}"
                            )
                            blip = pic.find(".//a:blip", nsmap)
                            if blip is not None:
                                embed_attr = blip.get(
                                    "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
                                )
                                if embed_attr:
                                    image_part = document.part.related_parts[embed_attr]
                                    image_data = BytesIO(image_part._blob)
                                    image = Image.open(image_data)
                                    image.load()
                                    block_images.append(image)
                                    image_counter += 1
            elif isinstance(block, Table):
                table_text = read_docx_tables(block)
                block_texts.append(table_text)
            if block_texts or block_images:
                block_text = "\n".join(block_texts).strip()
                if block_text or block_images:
                    chunks.append(
                        Chunk(path=file_path, text=block_text, images=block_images)
                    )
    except Exception as e:
        raise ValueError(f"Error processing DOCX file {file_path}: {e}")
    return chunks

def scrape_pptx(
    file_path: str,
    verbose: bool = False,
    include_output_images: bool = True,
    options: Optional[Dict[str, Any]] = None
) -> List[Chunk]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.shapes.picture import Picture
    from pptx.shapes.autoshape import Shape as AutoShape

    prs = Presentation(file_path)
    chunks = []
    # iterate through each slide in the presentation
    for slide in prs.slides:
        slide_texts = []
        slide_images = []
        # iterate through each shape in the slide
        for shape in slide.shapes:
            if shape.has_text_frame:
                auto_shape = cast(AutoShape, shape)
                for paragraph in auto_shape.text_frame.paragraphs:
                    text = paragraph.text
                    if len(slide_texts) == 0:
                        text = "# " + text  # header for first text of a slide
                    slide_texts.append(text)
            # extract images from shapes
            if include_output_images and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                pic = cast(Picture, shape)
                image_data = pic.image.blob
                image = Image.open(BytesIO(image_data))
                slide_images.append(image)
        # add slide to chunks if it has text or images
        if slide_texts or slide_images:
            text = "\n".join(slide_texts).strip()
            if not include_output_images:
                slide_images = []
            chunks.append(Chunk(path=file_path, text=text, images=slide_images))
    # return all chunks
    return chunks

def scrape_ipynb(
    file_path: str,
    verbose: bool = False,
    include_output_images: bool = True,
) -> List[Chunk]:
    with open(file_path, "r", encoding="utf-8") as file:
        notebook = json.load(file)
    chunks = []
    # parse cells in the notebook
    for cell in notebook["cells"]:
        texts = []
        images: List[Image.Image] = []
        cell_type = cell["cell_type"]
        # parse cell content based on type
        if verbose:
            print(f"[thepipe] Scraping cell {cell_type} from {file_path}")
        if cell_type == "markdown":
            text = "".join(cell["source"])
            if include_output_images:
                images = get_images_from_markdown(text)
            texts.append(text)
        elif cell_type == "code":
            source = "".join(cell["source"])
            texts.append(source)
            output_texts = []
            # code cells can have outputs
            if "outputs" in cell:
                for output in cell["outputs"]:
                    if (
                        include_output_images
                        and "data" in output
                        and "image/png" in output["data"]
                    ):
                        image_data = output["data"]["image/png"]
                        image = Image.open(BytesIO(base64.b64decode(image_data)))
                        images.append(image)
                    elif "data" in output and "text/plain" in output["data"]:
                        output_text = "".join(output["data"]["text/plain"])
                        output_texts.append(output_text)
            if output_texts:
                texts.extend(output_texts)
        elif cell_type == "raw":
            text = "".join(cell["source"])
            texts.append(text)
        if texts or images:
            text = "\n".join(texts).strip()
            chunks.append(Chunk(path=file_path, text=text, images=images))
    return chunks

def scrape_tweet(url: str, include_output_images: bool = True, verbose: bool = False, text_only: bool = False, options: Optional[Dict[str, Any]] = None) -> List[Chunk]:
    """
    Magic function from https://github.com/vercel/react-tweet/blob/main/packages/react-tweet/src/api/fetch-tweet.ts
    unofficial, could break at any time
    """

    def get_token(id: str) -> str:
        result = (float(id) / 1e15) * math.pi
        base_36_result = ""
        characters = "0123456789abcdefghijklmnopqrstuvwxyz"
        while result > 0:
            remainder = int(result % (6**2))
            base_36_result = characters[remainder] + base_36_result
            result = (result - remainder) // (6**2)
        base_36_result = re.sub(r"(0+|\.)", "", base_36_result)
        return base_36_result

    tweet_id = url.split("status/")[-1].split("?")[0]
    token = get_token(tweet_id)
    tweet_api_url = "https://cdn.syndication.twimg.com/tweet-result"
    params = {"id": tweet_id, "language": "en", "token": token}
    response = requests.get(tweet_api_url, params=params)
    if response.status_code != 200:
        raise ValueError(f"Failed to fetch tweet. Status code: {response.status_code}")
    tweet_data = response.json()
    # Extract tweet text
    tweet_text = tweet_data.get("text", "")
    # Extract images from tweet
    images: List[Image.Image] = []
    if include_output_images and "mediaDetails" in tweet_data:
        for media in tweet_data["mediaDetails"]:
            image_url = media.get("media_url_https")
            if image_url:
                image_response = requests.get(image_url)
                img = Image.open(BytesIO(image_response.content))
                images.append(img)
    # Create chunks for text and images
    chunk = Chunk(path=url, text=tweet_text, images=images)
    return [chunk]

def scrape_database(
    filepath: str,
    query: Optional[str] = None,
    db_type: Optional[str] = None,
    verbose: bool = False,
    options: Optional[Dict[str, Any]] = None
) -> List[Chunk]:
    """
    Scrape content from a database connection or file.
    """
    # Database operations are always local
    if verbose:
        print("[thepipe] Processing database locally")
    
    # Import database utilities
    from .database_utils import process_database
    
    # Determine mode from options
    mode = None
    if options:
        if options.get("schema_only"):
            mode = "schema"
        elif options.get("preview"):
            mode = "preview"
    
    # Process the database
    return process_database(
        connection_info=filepath,
        query=query,
        db_type=db_type,
        mode=mode,
        verbose=verbose,
        options=options
    )